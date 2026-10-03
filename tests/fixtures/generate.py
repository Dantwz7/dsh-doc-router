#!/usr/bin/env python3
"""Generate the synthetic fixtures used by the test suite.

Everything this script writes is **original synthetic content owned by this
repository** — no published paper, no third-party document, no personal data.
That is deliberate: the test suite must be reproducible and redistributable,
and real papers cannot be committed.

Usage
-----
    python tests/fixtures/generate.py            # write fixtures + self-check
    python tests/fixtures/generate.py --no-check # write only

The script imports the plugin's own probe (`lib/docprobe.py`) and asserts that
each generated PDF classifies the way the tests expect. If a future change to
the column-geometry heuristics breaks that, regenerating fixtures fails loudly
instead of silently producing fixtures that no longer test what they claim.

Requirements: PyMuPDF (required for the PDFs). python-docx / openpyxl /
python-pptx are optional — without them the corresponding Office fixtures are
skipped and a warning is printed; the committed copies stay valid.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from pathlib import Path

# This script imports `docprobe` for its self-check, and importing a module
# writes a __pycache__ beside it — which then gets picked up by the npm `files`
# allow-list and shipped in the tarball. Running the probe as a script never
# does this; only the import here would.
sys.dont_write_bytecode = True

HERE = Path(__file__).resolve().parent
REPO_ROOT = HERE.parent.parent
sys.path.insert(0, str(REPO_ROOT / "lib"))

PAGE_W, PAGE_H = 595.0, 842.0  # A4 in points
MARGIN = 56.0
TOP, BOTTOM = 64.0, 782.0
GUTTER = 30.0
FONT, SIZE, LEADING = "helv", 9.5, 12.0

# Mirrors `MIN_BLOCK_CHARS` in `lib/docprobe.py`, which drops any text block shorter
# than this before looking for gutters. Kept as a local copy so the fixtures stay
# importable without docprobe; `self_check()` asserts the two still agree, so a
# change on either side fails loudly instead of silently invalidating the fixture
# that exists to exercise the filter.
#
# The unit matters: since 0.5.0 this counts **latin-equivalent codepoints**, not raw
# codepoints (see `DENSE_CHAR_WEIGHT`). A pure-Latin block is unaffected, which is
# why `SHORT_BLOCK_CHARS` still describes the Latin fixtures exactly.
SHORT_BLOCK_CHARS = 25

# Mirrors `CJK_CHAR_WEIGHT`: one ideograph counts as this many latin-equivalent
# codepoints, because a Chinese word is 1-2 ideographs while an English word is
# ~5 letters. `self_check()` asserts this too.
DENSE_CHAR_WEIGHT = 2

# PyMuPDF's built-in Simplified-Chinese font. Deliberately **not** a system font:
# the fixture must regenerate identically on a machine with no CJK font installed.
# Verified to round-trip through `get_text()` — which is the whole point, since the
# block filter reads the extracted text.
DENSE_FONT = "china-s"

# Right-column entries for `dense-script-column.pdf`. Every one is short *as
# codepoints* yet substantial *as content* — the exact disagreement the weighted
# threshold exists to fix. Both halves are asserted in the builder.
DENSE_RIGHT_ENTRIES = (
    "界面残留会显著降低器件性能",
    "环境稳定性决定工艺窗口宽窄",
    "晶格匹配程度影响外延薄膜质量",
    "水溶性牺牲层可在数分钟内溶解",
    "工艺兼容性要求可原位生长薄膜",
    "衬底可重复使用且物性不退化",
    "远程外延依赖衬底表面的电荷分布",
    "二维材料转移需要完整的界面接触",
    "牺牲层厚度影响剥离后的表面粗糙度",
    "化学腐蚀速率必须与薄膜厚度匹配",
    "外延层与衬底之间保持共格界面",
    "退火温度决定结晶质量与缺陷密度",
    "大面积制备仍需解决均匀性问题",
    "器件性能对界面缺陷非常敏感",
)

# Synthetic stand-in for the publisher stamp that runs down a PDF's page edge.
MARGIN_STAMP = (
    "15214095, 2024, 15, Downloaded from https://example.invalid/library "
    "by a subscribing institution on 01 October 2026. "
) * 6

SENTENCES = (    "Document routing decides which extraction pipeline a file deserves.",
    "A converter that ignores page geometry will interleave columns of text.",
    "The probe reads the file header instead of trusting the extension.",
    "Layout is measured per horizontal band, not from the page as a whole.",
    "A scanned page carries no text layer and must be rendered as an image.",
    "Structure is lost when a layout-aware reader is replaced by a naive one.",
    "Magic bytes are authoritative because extensions are frequently wrong.",
    "Every decision here is deterministic, offline, and free of model calls.",
    "The gutter between columns is found by merging text intervals per band.",
    "Superscripts, italics and headings survive a layout-aware conversion.",
    "A single reference page must not outvote multi-column body pages.",
    "The correct pipeline for a two-column paper is the layout-aware one.",
)


def paragraphs(rng: random.Random, count: int) -> list[str]:
    """Build `count` paragraphs of plausible filler prose."""
    return [
        " ".join(rng.choice(SENTENCES) for _ in range(rng.randint(3, 5)))
        for _ in range(count)
    ]


def wrap(text: str, width_chars: int) -> list[str]:
    """Greedy word wrap. Helvetica's average glyph is ~0.5em wide, which is
    accurate enough to keep lines inside their column without a font metric
    round trip."""
    lines: list[str] = []
    current = ""
    for word in text.split():
        candidate = f"{current} {word}".strip()
        if current and len(candidate) > width_chars:
            lines.append(current)
            current = word
        else:
            current = candidate
    if current:
        lines.append(current)
    return lines


def _draw_column(
    page,
    x: float,
    width: float,
    lines: list[str],
    gap_lines: int,
    *,
    fontname: str = FONT,
    char_width_factor: float = 0.5,
) -> None:
    """Draw wrapped lines down a column, leaving `gap_lines` blank lines between
    paragraphs so the extractor sees them as distinct blocks.

    `char_width_factor` is the average glyph advance as a fraction of the font
    size: ~0.5em for Helvetica, but ~1.0em for CJK, where every glyph is full
    width. Both defaults reproduce the original Latin-only behaviour, so the
    pre-existing fixtures regenerate byte-identically.
    """
    y = TOP
    width_chars = max(20, int(width / (char_width_factor * SIZE)))
    for text in lines:
        for line in wrap(text, width_chars):
            if y > BOTTOM:
                return
            page.insert_text((x, y), line, fontname=fontname, fontsize=SIZE)
            y += LEADING
        y += gap_lines * LEADING


def build_pdf(path: Path, columns: int, seed: int, gap_lines: int) -> None:
    """Render a one- or two-column text page. Column count is real geometry:
    two text boxes side by side, not two columns faked with whitespace."""
    import pymupdf

    rng = random.Random(seed)
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)

    if columns == 1:
        _draw_column(page, MARGIN, PAGE_W - 2 * MARGIN, paragraphs(rng, 10), gap_lines)
    else:
        col_w = (PAGE_W - 2 * MARGIN - GUTTER) / 2
        _draw_column(page, MARGIN, col_w, paragraphs(rng, 6), gap_lines)
        _draw_column(page, MARGIN + col_w + GUTTER, col_w, paragraphs(rng, 6), gap_lines)

    doc.save(str(path))
    doc.close()


def build_reading_order_pdf(path: Path, marker_count: int) -> None:
    """A two-column page whose columns are *individually identifiable*.

    `two-column.pdf` fills both columns from one shared sentence pool, so every
    permutation of its sentences reads as plausibly as the correct one — it can
    verify the column *classification*, but it cannot verify the *reading order*.
    This fixture closes that hole: the left column carries only `L##` markers and
    the right column only `R##`, so the correct output is exactly
    `L01…Lnn` followed by `R01…Rnn`, and any interleaving, column swap or dropped
    column changes the sequence.

    Same geometry as `two-column.pdf` (identical margins, gutter and font), so
    the column heuristic under test is exercised exactly as it is there.

    Each marker is its own paragraph, separated by a blank line. That is not
    cosmetic: PyMuPDF merges vertically contiguous lines into a *single* text
    block, and `_column_profile` needs at least six blocks on the page before it
    will report anything at all. With no inter-paragraph gap this fixture routes
    as `unknown` — verified, which is why `gap_lines` is 1 here.
    """
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    col_w = (PAGE_W - 2 * MARGIN - GUTTER) / 2
    left = [f"L{i:02d} left column reading order marker." for i in range(1, marker_count + 1)]
    right = [f"R{i:02d} right column reading order marker." for i in range(1, marker_count + 1)]
    _draw_column(page, MARGIN, col_w, left, 1)
    _draw_column(page, MARGIN + col_w + GUTTER, col_w, right, 1)

    doc.save(str(path))
    doc.close()


def build_three_column_pdf(path: Path, marker_count: int) -> None:
    """Three side-by-side columns, each individually identifiable.

    Covers the branch `two-column-reading-order.pdf` cannot: `_band_columns`
    counting **three** columns. The rule under test is "count the columns", not
    "find the gutter down the middle" — a middle-column layout is exactly what the
    original middle-gutter implementation missed (see the routing-history notes in
    the README).

    Two hard constraints, both enforced by the assertion below, because a third
    column is only ~141pt wide here:

    - Every marker must be **at least 25 characters**. `_column_profile` drops any
      text block shorter than that, so a 24-character marker is invisible to the
      column detector: the page then has only the middle column's blocks left and
      routes as `single-column`. Verified the hard way.
    - Every marker must fit **one line** (`width_chars`), or it wraps and the
      paragraph takes two lines of height.

    Like the two-column fixture, each marker is its own paragraph so the page has
    the six-plus text blocks `_column_profile` requires.
    """
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    col_w = (PAGE_W - 2 * MARGIN - 2 * GUTTER) / 3
    width_chars = max(20, int(col_w / (0.5 * SIZE)))

    for index, (tag, name) in enumerate((("L", "leftmost"), ("M", "middle"), ("R", "rightmost"))):
        lines = [f"{tag}{i:02d} {name} column marker." for i in range(1, marker_count + 1)]
        for line in lines:
            assert len(line) >= 25, f"block filter would drop {line!r} ({len(line)} chars)"
            assert len(line) <= width_chars, f"{line!r} wraps in {width_chars} chars"
        x = MARGIN + index * (col_w + GUTTER)
        _draw_column(page, x, col_w, lines, 1)

    doc.save(str(path))
    doc.close()


def build_short_column_pdf(path: Path) -> None:
    """Two columns where the **right** one is entirely short blocks.

    Regression for the 25-character block filter. `_column_profile` drops every
    text block shorter than `SHORT_BLOCK_CHARS` before it looks for gutters, and
    that filter is column-agnostic: when one column is a figure, a table, or a list
    of short entries, it contributes no evidence at all. The surviving column then
    supplies more than the six blocks and eight bands the detector needs, so the
    page looked like a **confident single-column** page and the document went to
    `markitdown` — the one direction the asymmetric bet says must never happen.

    The expected verdict is `unknown`, not `multi-column`. The detector has no
    evidence for multi-column here, and asserting it would be a guess in the other
    direction; what the fixture pins is the *safe* outcome — evidence incomplete,
    so route to `pdf_markdown`.

    Both halves of the premise are asserted, so a font or geometry change cannot
    quietly turn this into an ordinary two-column fixture that tests nothing.
    """
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    col_w = (PAGE_W - 2 * MARGIN - GUTTER) / 2

    # Left column: ordinary prose, every line comfortably over the threshold.
    left = list(SENTENCES)
    for line in left:
        assert len(line) >= SHORT_BLOCK_CHARS, f"kept line is too short: {line!r}"

    # Right column: a figure/table-ish list, every entry under the threshold.
    right = [
        "Fig. 3a",
        "Fig. 3b",
        "Table 2",
        "n = 41",
        "P < .05",
        "et al.",
        "12.5%",
        "see Fig. 1",
        "20 um",
        "RT",
        "~3 eV",
        "x 10^4",
    ]
    for line in right:
        assert len(line) < SHORT_BLOCK_CHARS, f"dropped line would be kept: {line!r}"

    _draw_column(page, MARGIN, col_w, left, 1)
    _draw_column(page, MARGIN + col_w + GUTTER, col_w, right, 1)

    doc.save(str(path))
    doc.close()


def build_dense_script_column_pdf(path: Path) -> None:
    """`short-column.pdf`'s geometry, but the short right column is **Chinese**.

    A controlled comparison, and the regression for the *unit* the block filter
    counts in. `short-column.pdf` holds the layout constant and varies nothing but
    which column is short; this one holds the layout *and the left column* constant
    and varies only the **script** of the short column.

    The threshold `MIN_BLOCK_CHARS` is a *content* threshold — roughly "half a line
    of body text". Counted in raw codepoints it is not the same amount of content in
    both scripts: a Chinese word is 1-2 ideographs while an English word is ~5
    letters, so 25 codepoints of Chinese is far more text than 25 codepoints of
    Latin. A codepoint-counting filter therefore discards real Chinese body text
    that it keeps in English — and when the discarded text is a whole column, the
    page loses all evidence of that column.

    Before the fix: the right column vanished, `_dropped_blocks_look_like_a_column`
    saw surviving evidence on one side only, and the page came back `unknown` (safe,
    but not *right* — it had the evidence and threw it away). After the fix the
    right column survives the filter, the gutter is measured, and the page is
    correctly `multi-column`.

    Note the deliberate contrast with `short-column.pdf`, which stays `unknown`:
    "Fig. 3a" and "12.5%" really are page furniture and must keep being discarded.
    The fix does not loosen the filter for Latin — it only stops the filter from
    mistaking *Chinese* body text for furniture.

    Both halves of the premise are asserted, so a font, weight or threshold change
    cannot quietly turn this into a fixture that tests nothing.
    """
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)
    col_w = (PAGE_W - 2 * MARGIN - GUTTER) / 2

    # Left column: ordinary Latin prose, identical to short-column.pdf, comfortably
    # over the threshold in either unit.
    left = list(SENTENCES)
    for line in left:
        assert len(line) >= SHORT_BLOCK_CHARS, f"kept line is too short: {line!r}"

    # Right column: short but *meaningful* Chinese entries — a list of real points,
    # not page furniture. Each is under 25 codepoints yet at least 25 equivalent
    # ones, which is exactly the disagreement under test.
    right = list(DENSE_RIGHT_ENTRIES)
    for line in right:
        assert len(line) < SHORT_BLOCK_CHARS, (
            f"the old codepoint filter would keep {line!r} ({len(line)} codepoints) "
            "— the fixture no longer exercises the unit bug"
        )
        assert len(line) * DENSE_CHAR_WEIGHT >= SHORT_BLOCK_CHARS, (
            f"the weighted filter would drop {line!r} "
            f"({len(line) * DENSE_CHAR_WEIGHT} equivalent codepoints)"
        )

    _draw_column(page, MARGIN, col_w, left, 1)
    # CJK glyphs are full-width (~1em), so the wrap budget halves.
    _draw_column(
        page,
        MARGIN + col_w + GUTTER,
        col_w,
        right,
        1,
        fontname=DENSE_FONT,
        char_width_factor=1.0,
    )

    doc.save(str(path))
    doc.close()


def build_watermarked_pdf(path: Path, seed: int, gap_lines: int) -> None:
    """Two real columns plus a full-height, ~7pt-wide margin stamp.

    Publishers (Wiley, ACS) print a rotated "Downloaded from ..." string down
    the page edge. It is narrow but spans the whole page height, so it lands in
    *every* horizontal band. Counting it as a column made the estimator report
    three columns for a two-column paper — on a real 74-paper corpus that hit 18
    files. Regression here silently inflates `column_estimate` for most
    published PDFs.
    """
    import pymupdf

    rng = random.Random(seed)
    doc = pymupdf.open()
    page = doc.new_page(width=PAGE_W, height=PAGE_H)

    col_w = (PAGE_W - 2 * MARGIN - GUTTER) / 2
    _draw_column(page, MARGIN, col_w, paragraphs(rng, 6), gap_lines)
    _draw_column(page, MARGIN + col_w + GUTTER, col_w, paragraphs(rng, 6), gap_lines)

    # Rotated, so a single text block runs the height of the page at the edge.
    page.insert_text(
        (PAGE_W - 14, PAGE_H - 24),
        MARGIN_STAMP,
        fontname=FONT,
        fontsize=5.5,
        rotate=90,
    )
    doc.save(str(path))
    doc.close()


def build_unknown_pdf(path: Path, pages: int = 3) -> None:
    """A text-bearing PDF whose pages carry no usable column evidence.

    Each page's text arrives as one large block, and `_column_profile` needs
    several blocks to measure bands, so it returns None for every page and the
    probe must answer `columns: unknown` rather than guess. That state is
    load-bearing: it routes to the safe side (`pdf_markdown`), because guessing
    "single column" here would send a possibly multi-column paper to a converter
    that shreds its body text.
    """
    import pymupdf

    body = " ".join(SENTENCES) * 6
    doc = pymupdf.open()
    for _ in range(pages):
        page = doc.new_page(width=PAGE_W, height=PAGE_H)
        page.insert_textbox(
            pymupdf.Rect(MARGIN + 16, TOP, PAGE_W - MARGIN - 16, BOTTOM),
            body,
            fontname=FONT,
            fontsize=SIZE,
        )
    doc.save(str(path))
    doc.close()


def build_scanned(path: Path, source: Path) -> None:
    """Wrap a rendered page in an image-only PDF: zero text layer, so the
    router must recommend the image pipeline.

    JPEG rather than PNG, because a fixture is committed to the repository and
    a lossless full-page raster costs over a megabyte. Real scans are JPEG or
    CCITT anyway, so this is also the more faithful shape.
    """
    import pymupdf

    src = pymupdf.open(str(source))
    pix = src[0].get_pixmap(dpi=100, colorspace=pymupdf.csGRAY)
    src.close()

    out = pymupdf.open()
    page = out.new_page(width=pix.width, height=pix.height)
    page.insert_image(page.rect, stream=pix.tobytes("jpg", jpg_quality=55))
    out.save(str(path), deflate=True, garbage=3)
    out.close()


def build_png(path: Path) -> None:
    """A small chart-ish raster, for the `read_image` branch."""
    import pymupdf

    doc = pymupdf.open()
    page = doc.new_page(width=320, height=200)
    page.insert_text((24, 40), "Measured accuracy", fontname="helv", fontsize=13)
    for i, value in enumerate((0.81, 0.87, 0.79)):
        page.draw_rect(
            pymupdf.Rect(40 + i * 70, 160 - value * 110, 90 + i * 70, 160),
            color=None,
            fill=(0.20, 0.45, 0.75),
        )
        page.insert_text(
            (48 + i * 70, 178), f"m{i + 1}", fontname="helv", fontsize=10
        )
    page.get_pixmap(dpi=96).save(str(path))
    doc.close()


def _write_text(path: Path, text: str, encoding: str = "utf-8") -> None:
    """Write LF line endings on every platform.

    `Path.write_text` translates `\\n` to the platform separator, so on Windows
    the fixtures would land as CRLF while `.gitattributes` normalises the
    repository to LF — leaving every Windows contributor with a permanently
    dirty working tree.
    """
    with path.open("w", encoding=encoding, newline="\n") as fh:
        fh.write(text)


def build_text_fixtures() -> None:
    """Plain-text-ish fixtures. No library needed, so these are always written."""
    _write_text(
        HERE / "sample.txt",
        "Plain text needs no conversion.\nSecond line for good measure.\n",
    )
    _write_text(
        HERE / "sample.csv",
        "model,params,score\nalpha,1.3B,0.81\nbeta,7B,0.87\n",
    )
    _write_text(
        HERE / "sample.json",
        json.dumps({"model": "alpha", "score": 0.81, "tags": ["routing"]}, indent=2) + "\n",
    )
    _write_text(
        HERE / "sample.html",
        "<!doctype html>\n<html><body>\n<h1>Routing report</h1>\n"
        "<p>Body paragraph.</p>\n<ul><li>one</li><li>two</li></ul>\n"
        "</body></html>\n",
    )
    # A minimal, *valid* RTF document: {\rtf1 ... } with an escape and a group.
    _write_text(
        HERE / "sample.rtf",
        r"{\rtf1\ansi\deff0{\fonttbl{\f0 Helvetica;}}\f0\fs20 "
        r"Routing report\par Second paragraph with an escape: \'e9.\par}",
        encoding="ascii",
    )


def build_magic_only_fixtures() -> None:
    """Header-only probes for the magic-byte branch.

    These are deliberately NOT complete documents. `doc_route` classifies by
    file header, so a header-only file is the honest minimal input for that
    rule — and it keeps a 100-byte fixture from being mistaken for real
    content. See tests/fixtures/README.md.
    """
    (HERE / "magic-ole2.doc").write_bytes(
        b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + b"\x00" * 504
    )
    import zipfile

    with zipfile.ZipFile(HERE / "magic-plain.zip", "w", zipfile.ZIP_DEFLATED) as zf:
        zf.writestr("readme.txt", "A plain zip container: not an Office file.\n")


def build_office_fixtures() -> list[str]:
    """Office fixtures via their native libraries. Optional: skip if missing."""
    written: list[str] = []

    try:
        from docx import Document
        from docx.shared import Pt  # noqa: F401  (kept for future styling)

        doc = Document()
        doc.add_heading("Routing report", level=1)
        doc.add_paragraph("Generated by tests/fixtures/generate.py.")
        doc.add_heading("Method", level=2)
        for item in ("header sniffing", "column geometry", "text-layer density"):
            doc.add_paragraph(item, style="List Bullet")
        table = doc.add_table(rows=3, cols=2)
        table.style = "Table Grid"
        for r, row in enumerate([["group", "score"], ["alpha", "0.81"], ["beta", "0.87"]]):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
        doc.save(str(HERE / "sample.docx"))
        written.append("sample.docx")
    except ImportError:
        print("warn: python-docx missing -> sample.docx not regenerated", file=sys.stderr)

    try:
        from openpyxl import Workbook

        wb = Workbook()
        ws = wb.active
        ws.title = "scores"
        ws.append(["model", "params", "score"])
        ws.append(["alpha", "1.3B", 0.81])
        ws.append(["beta", "7B", 0.87])
        wb.save(str(HERE / "sample.xlsx"))
        written.append("sample.xlsx")
    except ImportError:
        print("warn: openpyxl missing -> sample.xlsx not regenerated", file=sys.stderr)

    try:
        from pptx import Presentation

        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[1])
        slide.shapes.title.text = "Routing report"
        slide.placeholders[1].text = "Format sniffing\nColumn geometry\nText-layer density"
        prs.save(str(HERE / "sample.pptx"))
        written.append("sample.pptx")
    except ImportError:
        print("warn: python-pptx missing -> sample.pptx not regenerated", file=sys.stderr)

    return written


EXPECTED = {
    "single-column.pdf": {"columns": "single-column", "text_layer": True},
    "two-column.pdf": {"columns": "multi-column", "text_layer": True, "column_estimate": 2},
    "two-column-reading-order.pdf": {
        "columns": "multi-column",
        "text_layer": True,
        "column_estimate": 2,
    },
    "three-column.pdf": {
        "columns": "multi-column",
        "text_layer": True,
        "column_estimate": 3,
    },
    "two-column-watermark.pdf": {
        "columns": "multi-column",
        "text_layer": True,
        "column_estimate": 2,
    },
    "unknown-columns.pdf": {
        "columns": "unknown",
        "text_layer": True,
        "recommend": ["pdf_markdown"],
    },
    "short-column.pdf": {
        "columns": "unknown",
        "text_layer": True,
        "recommend": ["pdf_markdown"],
    },
    "dense-script-column.pdf": {
        "columns": "multi-column",
        "text_layer": True,
        "column_estimate": 2,
        "recommend": ["pdf_markdown"],
    },
    "scanned.pdf": {"text_layer": False},
    "sample.docx": {"format": "docx"},
    "sample.xlsx": {"format": "xlsx"},
    "sample.pptx": {"format": "pptx"},
    "sample.png": {"format": "png"},
    "sample.rtf": {"format": "rtf"},
    "magic-ole2.doc": {"format": "ole2"},
    "magic-plain.zip": {"format": "zip"},
}


def self_check() -> int:
    """Assert each fixture routes the way the tests assume."""
    import docprobe  # noqa: E402  (path inserted above)

    failures: list[str] = []
    # The short-block fixture is only a fixture while these two numbers agree; if
    # docprobe's threshold moves, the right column would start being counted and the
    # fixture would quietly stop exercising the filter it was built for.
    if docprobe.MIN_BLOCK_CHARS != SHORT_BLOCK_CHARS:
        failures.append(
            f"MIN_BLOCK_CHARS drifted: docprobe has {docprobe.MIN_BLOCK_CHARS}, "
            f"generate.py assumes {SHORT_BLOCK_CHARS} — short-column.pdf is now stale"
        )
    # Same contract for the weighted unit: if the ideograph weight moves,
    # `dense-script-column.pdf`'s right column stops being "short by codepoints but
    # long by content" and the fixture silently stops testing the unit bug.
    if getattr(docprobe, "CJK_CHAR_WEIGHT", None) != DENSE_CHAR_WEIGHT:
        failures.append(
            f"CJK_CHAR_WEIGHT drifted: docprobe has "
            f"{getattr(docprobe, 'CJK_CHAR_WEIGHT', None)}, generate.py assumes "
            f"{DENSE_CHAR_WEIGHT} — dense-script-column.pdf is now stale"
        )
    for name, expected in EXPECTED.items():
        path = HERE / name
        if not path.exists():
            failures.append(f"{name}: missing")
            continue
        result = docprobe.route(path)
        for key, want in expected.items():
            got = result.get(key)
            if got != want:
                failures.append(f"{name}: {key} expected {want!r}, got {got!r}")
        detail = {
            k: result.get(k)
            for k in ("format", "columns", "column_estimate", "text_layer", "chars_per_page")
            if k in result
        }
        print(f"  {name:20s} {detail}")

    if failures:
        print("\nSELF-CHECK FAILED:", file=sys.stderr)
        for line in failures:
            print(f"  - {line}", file=sys.stderr)
        return 1
    print("\nself-check passed: every fixture routes as expected")
    return 0


def main(argv: list[str]) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument(
        "--no-check",
        action="store_true",
        help="write fixtures without asserting their routing classification",
    )
    parser.add_argument(
        "--gap-lines",
        type=int,
        default=1,
        help="blank lines between paragraphs in the generated PDFs (default 1)",
    )
    args = parser.parse_args(argv)

    try:
        import pymupdf  # noqa: F401
    except ImportError:
        print(
            "error: PyMuPDF is required to generate the PDF fixtures.\n"
            "       pip install pymupdf4llm",
            file=sys.stderr,
        )
        return 2

    HERE.mkdir(parents=True, exist_ok=True)
    build_pdf(HERE / "single-column.pdf", columns=1, seed=7, gap_lines=args.gap_lines)
    build_pdf(HERE / "two-column.pdf", columns=2, seed=11, gap_lines=args.gap_lines)
    build_reading_order_pdf(HERE / "two-column-reading-order.pdf", marker_count=15)
    build_three_column_pdf(HERE / "three-column.pdf", marker_count=15)
    build_short_column_pdf(HERE / "short-column.pdf")
    build_dense_script_column_pdf(HERE / "dense-script-column.pdf")
    build_watermarked_pdf(HERE / "two-column-watermark.pdf", seed=23, gap_lines=args.gap_lines)
    build_unknown_pdf(HERE / "unknown-columns.pdf")
    build_scanned(HERE / "scanned.pdf", HERE / "single-column.pdf")
    build_png(HERE / "sample.png")
    build_text_fixtures()
    build_magic_only_fixtures()
    office = build_office_fixtures()

    print(f"wrote fixtures to {HERE}")
    if office:
        print("office fixtures:", ", ".join(office))
    if args.no_check:
        return 0
    print("\nrouting self-check:")
    return self_check()


if __name__ == "__main__":
    raise SystemExit(main(sys.argv[1:]))
